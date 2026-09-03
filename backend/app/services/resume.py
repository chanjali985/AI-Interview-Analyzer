"""Resume ingestion: turn an uploaded PDF/DOCX/TXT into plain text."""
from __future__ import annotations

import logging
import os
import re

logger = logging.getLogger(__name__)

MAX_RESUME_CHARS = 20_000


class ResumeParseError(RuntimeError):
    """Raised when a resume file cannot be read."""


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_RESUME_CHARS]


def _from_pdf(path: str) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ResumeParseError("pypdf is not installed; cannot read PDF resumes") from exc

    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:  # noqa: BLE001
            raise ResumeParseError("This PDF is password protected") from exc
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001
            logger.debug("Failed to extract a PDF page: %s", exc)
    return "\n".join(parts)


def _from_docx(path: str) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise ResumeParseError("python-docx is not installed; cannot read .docx resumes") from exc

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _from_txt(path: str) -> str:
    with open(path, encoding="utf-8", errors="replace") as handle:
        return handle.read()


def extract_text(path: str) -> str:
    """Extract plain text from a resume file."""
    if not os.path.exists(path):
        raise ResumeParseError(f"Resume file not found: {path}")

    extension = os.path.splitext(path)[1].lower()
    if extension == ".pdf":
        text = _from_pdf(path)
    elif extension == ".docx":
        text = _from_docx(path)
    elif extension in (".txt", ".md", ""):
        text = _from_txt(path)
    else:
        raise ResumeParseError(f"Unsupported resume format: {extension}")

    cleaned = _clean(text)
    if len(cleaned) < 40:
        raise ResumeParseError(
            "Could not read meaningful text from this resume. "
            "If it is a scanned image, please upload a text-based PDF or paste the text instead."
        )
    return cleaned
